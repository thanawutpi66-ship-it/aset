from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path("output/docx/chapter_5_verification_validation_draft.docx")
FONT = "TH Sarabun New"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="B7B7B7", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_text(doc, text, *, bold=False, italic=False, align=None, after=4, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, 16, bold)
    r.italic = italic
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run_font(r, 15)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = "Heading %d" % min(level + 1, 3)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, 18 if level == 1 else 16, True)
    return p


def format_table(table, widths=None):
    table.style = "Table Grid"
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            if row_index == 0:
                set_cell_shading(cell, "1F4E78")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F2F6FA")
            if widths:
                cell.width = Cm(widths[col_index])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_run_font(run, 12.5, row_index == 0, (255, 255, 255) if row_index == 0 else None)
                if row_index == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_table(doc, caption, headers, rows, widths=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(caption)
    set_run_font(r, 14, True)
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = str(text)
    format_table(table, widths)
    return table


def add_figure_placeholder(doc, number, title, instruction, height_cm=4.2):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, "F3F6F8")
    set_cell_border(cell, "7F8C8D", "10")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell.height = Cm(height_cm)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("พื้นที่สำหรับรูปที่ %s" % number)
    set_run_font(r, 16, True, (31, 78, 120))
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(instruction)
    set_run_font(r2, 12, False, (80, 80, 80))
    cap = doc.add_paragraph()
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(7)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = cap.add_run("รูปที่ %s %s" % (number, title))
    set_run_font(r3, 13, True)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.54)
    sec.right_margin = Cm(2.54)
    add_page_number(sec)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(16)
    for style_name, size in (("Heading 1", 20), ("Heading 2", 18), ("Heading 3", 16)):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("บทที่ 5 การพิสูจน์ยืนยันและการตรวจสอบความถูกต้อง\nChapter 5: Verification and Validation")
    set_run_font(r, 20, True)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_after = Pt(12)
    rn = note.add_run("ร่างเนื้อหาสำหรับบันทึกผลการทดสอบจริง — ตารางและรูปที่มีคำว่า [ใส่ผล] ต้องเติมจากการทดสอบรอบสุดท้าย")
    set_run_font(rn, 13, False, (117, 117, 117))

    add_text(doc, "บทนี้พิสูจน์ว่าระบบทดสอบสมรรถนะและประเมินสุขภาพแบตเตอรี่ (ASET Battery Tester) บรรลุข้อกำหนดทางวิศวกรรมในบทที่ 2 หรือไม่ โดยแยกผลที่ผ่าน ผลที่ไม่ผ่าน และผลที่ยังต้องยืนยันด้วยการทดสอบอ้างอิงอย่างชัดเจน การทดสอบครอบคลุมความถูกต้องของการวัด ความปลอดภัย การทำงานอัตโนมัติ ความน่าเชื่อถือของอัลกอริทึม และความครบถ้วนของข้อมูลตั้งแต่เครื่องทดสอบจนถึงแดชบอร์ด")

    add_heading(doc, "5.1 แผนการพิสูจน์ยืนยัน / Verification Plan")
    add_text(doc, "การพิสูจน์ยืนยันใช้หลักการเชื่อมโยงข้อกำหนดในบทที่ 2 กับวิธีทดสอบ เกณฑ์ยอมรับ และหลักฐานที่ตรวจสอบย้อนหลังได้ ตารางที่ 5.1 แสดงแผนการทดสอบหลักของระบบ")
    add_table(doc, "ตารางที่ 5.1 แผนการพิสูจน์ยืนยันข้อกำหนดของระบบ", ["กลุ่ม", "ข้อกำหนดที่พิสูจน์", "วิธีตรวจสอบ", "หลักฐาน/เกณฑ์"], [
        ["1", "Measurement & logging", "เทียบแรงดัน/กระแสกับเครื่องมืออ้างอิง และวิเคราะห์ Timestamp แยกตาม Mode", "ตาราง error, median/p95 Δt, CSV ดิบ"],
        ["2", "Safety & interlock", "จำลอง UVP, OVP, OTP และ E-stop อย่างละ 3 รอบ", "response time, waveform, alarm log"],
        ["3", "Sequence automation", "รัน Quick Scan, HPPC, charge/discharge แบบ end-to-end", "phase/Mode, completion status, CSV ครบ"],
        ["4", "Battery algorithm", "ทดสอบซ้ำและเทียบ Quick Scan กับ HPPC และ C10 reference", "R₀/R₁/C₁/τ/R², %CV, SoC/SoH error"],
        ["5", "Data & cloud", "ตรวจ metadata, CSV, cloud record และรายงานจาก Run ID เดียวกัน", "traceability ครบทุกปลายทาง"],
    ], [1.1, 3.7, 6.1, 6.2])

    add_heading(doc, "5.2 ขั้นตอนการทดสอบ / Testing Procedures")
    add_text(doc, "เพื่อให้เปรียบเทียบผลได้อย่างเป็นธรรม แบตเตอรี่ทุกลูกต้องระบุ Battery ID, profile, rated capacity, สภาพใช้งาน, อุณหภูมิห้อง, รุ่นซอฟต์แวร์ และ Run ID ก่อนเริ่มทดสอบ หากเป็นการเปรียบเทียบความจุหรือ SoH ต้องชาร์จเต็มและพักแบตตามเงื่อนไขที่กำหนดก่อนทุก run")
    add_table(doc, "ตารางที่ 5.2 เงื่อนไขทดสอบและเครื่องมือที่ต้องบันทึก", ["รายการ", "เงื่อนไข/เครื่องมือ", "ค่าที่ต้องบันทึก"], [
        ["แบตเตอรี่", "อย่างน้อย 3 สภาพ: ดี, ใช้งานปกติ, เสื่อม; ใช้ profile ที่ถูกต้อง", "Battery ID, รุ่น, rated Ah, อายุ/สภาพ"],
        ["สภาพแวดล้อม", "ห้องปฏิบัติการ; บันทึกทุก run", "อุณหภูมิห้อง, อุณหภูมิแบต"],
        ["เครื่องมืออ้างอิง", "DMM/oscilloscope ที่ระบุรุ่นและสถานะสอบเทียบ", "รุ่นเครื่องมือ, วันสอบเทียบ"],
        ["ข้อมูลทดสอบ", "CSV และ cloud record จาก Run ID เดียวกัน", "Mode, profile, app version, timestamp"],
        ["การทดสอบ reference", "HPPC สำหรับ dynamic/DCIR และ C10 สำหรับ capacity", "test current, cutoff, charge/rest condition"],
    ], [3.0, 8.0, 6.1])

    add_text(doc, "(1) Measurement & Logging: ตั้งค่าแรงดันและกระแสอย่างน้อย 5 จุดในช่วงใช้งานจริง วัดด้วยเครื่องมืออ้างอิงและเปรียบเทียบกับค่าที่ระบบบันทึก จากนั้นเก็บ Timestamp อย่างน้อย 1,000 samples ต่อ mode ที่สำคัญ เพื่อนำมาคำนวณ median Δt, p95 Δt, ค่าสูงสุด และจำนวนข้อมูลที่เกิน 0.5 s")
    add_text(doc, "(2) Safety & Interlock: ทดสอบ UVP, OVP, OTP และ E-stop อย่างละ 3 รอบ วัดเวลาตั้งแต่เกิดเงื่อนไข fault จนกระทั่งกระแสโหลดลดลงเป็นศูนย์ และเก็บ waveform, alarm log และสถานะ SSR/circuit breaker เป็นหลักฐาน")
    add_text(doc, "(3) Sequence Automation: รัน Quick Scan, HPPC และ charge/discharge ให้ครบทุก phase ตรวจว่ามี Mode ครบ, ไม่มี phase ข้าม, CSV ปิดไฟล์ได้ถูกต้อง และระบบแจ้งสถานะจบการทดสอบอย่างเหมาะสม")
    add_text(doc, "(4) Battery Algorithm: สำหรับแบตแต่ละลูก รัน Quick Scan อย่างน้อย 3 รอบ โดยชาร์จและพักให้มีเงื่อนไขเท่ากันทุกครั้ง เปรียบเทียบ DCIR/ECM กับ HPPC และเปรียบเทียบ SoH กับ C10 reference test ซึ่งใช้กระแสอ้างอิง C10 และ cutoff 1.75 V/cell")
    add_text(doc, "(5) Data & Cloud: ใช้ Run ID เดียวกันตรวจสอบความต่อเนื่องของข้อมูลตั้งแต่ CSV metadata, results analysis, cloud dashboard จนถึงรายงานที่สร้างจากผลทดสอบ")
    add_figure_placeholder(doc, "5.1", "ลำดับการทดสอบและความสัมพันธ์ของ Quick Scan, HPPC และ C10 reference", "วาด flowchart: Calibration → Safety → HPPC → Quick Scan (3 รอบ) → C10 reference → Grade validation. ใช้ลูกศรแสดงว่า HPPC เป็น reference ด้าน dynamic/DCIR และ C10 เป็น reference ด้าน capacity/SoH.", 4.6)

    add_heading(doc, "5.3 ผลการทดสอบ / Test Results")
    add_heading(doc, "5.3.1 ความถูกต้องของการวัดและ Adaptive Data Logging", level=2)
    add_text(doc, "ผลการทดสอบส่วนนี้ต้องรายงานความคลาดเคลื่อนของแรงดันและกระแสเทียบกับเครื่องมืออ้างอิง รวมถึงสถิติ Timestamp แยกตาม Mode เพื่อพิสูจน์ว่าระบบใช้ adaptive data logging จริง ไม่ใช่อ้างอัตราการสุ่มตัวอย่างเพียงค่าเดียวทั้งไฟล์")
    add_table(doc, "ตารางที่ 5.3 ความถูกต้องของแรงดัน กระแส และการตรวจสอบ Kelvin sensing", ["จุดทดสอบ", "Reference", "ASET measured", "Error", "%Error", "Pass/Fail"], [
        ["Voltage 1", "[ใส่ผล] V", "[ใส่ผล] V", "[ใส่ผล] V", "[ใส่ผล]", "[ใส่ผล]"],
        ["Voltage 2–5", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
        ["Current 1", "[ใส่ผล] A", "[ใส่ผล] A", "[ใส่ผล] A", "[ใส่ผล]", "[ใส่ผล]"],
        ["Current 2–5", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
        ["Kelvin comparison", "Force-side V [ใส่ผล]", "Sense-side V [ใส่ผล]", "ΔV [ใส่ผล]", "—", "[ใส่ผล]"],
    ], [2.6, 3.2, 3.2, 2.5, 2.2, 2.4])
    add_figure_placeholder(doc, "5.2", "กราฟความถูกต้องของแรงดันและกระแสเทียบเครื่องมืออ้างอิง", "สร้างกราฟ 2 panel: แรงดันและกระแส. แกน x = reference, แกน y = ASET measured; ใส่เส้น y=x และ error bar. ห้ามใช้ screenshot ตารางแทนกราฟ.", 4.2)
    add_table(doc, "ตารางที่ 5.4 สถิติ Adaptive Data Logging แยกตาม phase", ["ไฟล์/Mode", "จำนวน interval", "Median Δt (s)", "p95 Δt (s)", "Max Δt (s)", "ผล"], [
        ["Quick 1 – mini-pulse*", "139", "0.171", "0.353", "[ดู raw]", "ผ่านเบื้องต้น (≈5.8 Hz)"],
        ["Quick 1 – relaxation*", "286", "0.290", "0.355", "[ดู raw]", "ผ่านเบื้องต้น (≈3.4 Hz)"],
        ["Quick 1 – main discharge*", "320", "5.065", "16.002", "[ดู raw]", "low-rate; ต้องใช้เฉพาะ trend"],
        ["Quick 2 – mini-pulse*", "154", "0.178", "0.345", "[ดู raw]", "ผ่านเบื้องต้น (≈5.6 Hz)"],
        ["Quick 2 – relaxation*", "315", "0.310", "0.359", "[ดู raw]", "ผ่านเบื้องต้น (≈3.2 Hz)"],
        ["Quick 2 – main discharge*", "473", "5.042", "5.074", "[ดู raw]", "low-rate; ต้องเพิ่มใกล้ cutoff"],
    ], [3.8, 2.3, 2.5, 2.3, 2.4, 4.2])
    add_text(doc, "หมายเหตุ: แถวที่มีเครื่องหมาย * เป็นการแบ่ง phase จากเวลาและกระแสของ CSV เดิม ซึ่งยังไม่มีคอลัมน์ Mode จึงเป็นผลเบื้องต้นเท่านั้น การทดสอบรอบสุดท้ายต้องใช้ CSV ที่มี Mode เพื่อยืนยันผลโดยตรง", italic=True, after=5)
    add_figure_placeholder(doc, "5.3", "การกระจายของช่วงเวลาระหว่างตัวอย่าง (Δt) แยกตาม Mode", "ใช้ boxplot หรือ histogram แยก MINI_PULSE, RELAX, MAIN_DISCHARGE และ NEAR_CUTOFF. แกน y เป็น Δt (s) และขีดเส้นเกณฑ์ 0.5 s. ใช้ CSV ใหม่ที่มี Mode.", 4.2)

    add_heading(doc, "5.3.2 ความปลอดภัยและ Hardware Interlock", level=2)
    add_text(doc, "ผล safety ต้องรายงานทั้งเวลาตอบสนองและค่าจริง ณ จุดตัด ไม่ควรรายงานเพียงคำว่า PASS โดยไม่มี waveform หรือ log ประกอบ")
    add_table(doc, "ตารางที่ 5.5 ผลการทดสอบ UVP, OVP, OTP และ E-stop", ["การทดสอบ", "รอบ 1", "รอบ 2", "รอบ 3", "ค่าเฉลี่ย/ค่าตัด", "Pass/Fail"], [
        ["UVP response time", "[ใส่ผล] ms", "[ใส่ผล] ms", "[ใส่ผล] ms", "[ใส่ผล] ms", "[ใส่ผล]"],
        ["UVP cutoff voltage", "[ใส่ผล] V", "[ใส่ผล] V", "[ใส่ผล] V", "target 10.4–10.6 V", "[ใส่ผล]"],
        ["OVP response time", "[ใส่ผล] ms", "[ใส่ผล] ms", "[ใส่ผล] ms", "[ใส่ผล] ms", "[ใส่ผล]"],
        ["OTP trip temperature", "[ใส่ผล] °C", "[ใส่ผล] °C", "[ใส่ผล] °C", "target 45–60 °C", "[ใส่ผล]"],
        ["E-stop response time", "[ใส่ผล] ms", "[ใส่ผล] ms", "[ใส่ผล] ms", "[ใส่ผล] ms", "[ใส่ผล]"],
    ], [3.2, 2.6, 2.6, 2.6, 3.5, 2.4])
    add_figure_placeholder(doc, "5.4", "Waveform การตอบสนองของ SSR ต่อเหตุการณ์ fault", "ภาพจาก oscilloscope อย่างน้อย 2 channels: channel 1 = voltage/alarm trigger, channel 2 = load current หรือ SSR output. ใส่ cursor แสดงเวลา fault ถึง current = 0.", 4.5)

    add_heading(doc, "5.3.3 การทำงานอัตโนมัติของลำดับทดสอบ", level=2)
    add_table(doc, "ตารางที่ 5.6 ความครบถ้วนของ sequence และไฟล์ข้อมูล", ["โหมด", "Phase ที่ต้องครบ", "CSV/Mode ครบ", "Alarm/stop ถูกต้อง", "ผล"], [
        ["Quick Scan", "OCV, MINI_PULSE, RELAX, MAIN_DISCHARGE, TAIL_REST", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
        ["HPPC", "REST, pulse, relaxation ทุก cycle", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
        ["Charge/Discharge", "start, control, cutoff, complete", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
    ], [3.0, 6.4, 3.4, 3.6, 2.3])
    add_figure_placeholder(doc, "5.5", "ตัวอย่าง timeline ของ Quick Scan และ HPPC", "สร้างกราฟ 2 ส่วน: (ก) Quick Scan และ (ข) HPPC. แกน x = เวลา; เส้นแรงดันและกระแส; แถบสีด้านล่างแสดง Mode. ใส่ตำแหน่ง pulse, relax, cutoff และ tail rest.", 5.0)

    add_heading(doc, "5.3.4 ความน่าเชื่อถือของอัลกอริทึมประเมินแบตเตอรี่", level=2)
    add_text(doc, "การตัดสินผลอัลกอริทึมต้องแยก Quick Electrical Grade ออกจาก Capacity Grade โดย Quick Electrical Grade ใช้ DCIR/ECM ที่ผ่าน data-quality gate ส่วน Verified Overall Grade ต้องเทียบกับ C10 reference test และเลือกผลที่อนุรักษ์นิยมกว่าเมื่อผลสองส่วนไม่สอดคล้องกัน")
    add_table(doc, "ตารางที่ 5.7 การทำซ้ำของ DCIR และ ECM จาก Quick Scan", ["Battery ID", "Run", "R₀ (mΩ)", "R₁ (mΩ)", "C₁ (F)", "τ (s)", "R²", "ผล"], [
        ["B1 – สภาพดี", "1", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
        ["B1 – สภาพดี", "2", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
        ["B1 – สภาพดี", "3", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
        ["B2/B3 – ทำรูปแบบเดียวกัน", "1–3", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
        ["Summary", "mean / SD / %CV", "[ใส่ผล]", "—", "—", "—", "—", "target R₀ CV ≤15%"],
    ], [3.1, 1.6, 2.0, 2.0, 2.0, 1.8, 1.5, 2.6])
    add_figure_placeholder(doc, "5.6", "การ fitting แบบจำลอง 1-RC Thevenin ECM", "กราฟเส้น measured voltage และ fitted voltage ซ้อนกันในช่วง mini-pulse/relaxation; แสดงค่า R₀, R₁, C₁, τ และ R² ใน legend. เลือกอย่างน้อยหนึ่งกราฟต่อสภาพแบต.", 4.6)
    add_figure_placeholder(doc, "5.7", "ความสามารถในการทำซ้ำของ DCIR", "ใช้ bar chart หรือ dot plot: แกน x = Battery ID, แกน y = R₀ (mΩ), แสดง 3 run และ error bar (mean ± SD). ใส่ค่า %CV เหนือกลุ่มข้อมูล.", 4.2)
    add_table(doc, "ตารางที่ 5.8 การเปรียบเทียบ Quick Scan กับ HPPC reference", ["Battery ID", "Quick R₀", "HPPC R₀", "Difference", "%Difference", "แนวโน้มสอดคล้อง"], [
        ["B1", "[ใส่ผล] mΩ", "[ใส่ผล] mΩ", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
        ["B2", "[ใส่ผล] mΩ", "[ใส่ผล] mΩ", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
        ["B3", "[ใส่ผล] mΩ", "[ใส่ผล] mΩ", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
    ], [2.4, 3.0, 3.0, 3.0, 3.0, 4.0])
    add_figure_placeholder(doc, "5.8", "ความสัมพันธ์ระหว่าง DCIR จาก Quick Scan และ HPPC", "Scatter plot: แกน x = HPPC R₀ (mΩ), แกน y = Quick Scan R₀ (mΩ), ใส่เส้น y=x. ใช้เพื่อแสดงแนวโน้ม ไม่อ้างว่าค่า DCIR และ ACIR เท่ากัน.", 4.2)
    add_table(doc, "ตารางที่ 5.9 การยืนยัน SoC/SoH กับ C10 reference test", ["Battery ID", "Quick SoH", "C10 capacity/SoH", "Absolute error", "Verified grade", "ผล"], [
        ["B1", "[ใส่ผล] %", "[ใส่ผล] %", "[ใส่ผล] pp", "[ใส่ผล]", "[ใส่ผล]"],
        ["B2", "[ใส่ผล] %", "[ใส่ผล] %", "[ใส่ผล] pp", "[ใส่ผล]", "[ใส่ผล]"],
        ["B3", "[ใส่ผล] %", "[ใส่ผล] %", "[ใส่ผล] pp", "[ใส่ผล]", "[ใส่ผล]"],
    ], [2.4, 2.8, 3.8, 3.1, 3.0, 3.2])
    add_table(doc, "ตารางที่ 5.10 การตรวจสอบผลการคัดเกรดแบตเตอรี่", ["Battery ID/สภาพ", "Quick Electrical Grade", "Capacity Grade", "Verified Overall Grade", "คำอธิบาย"], [
        ["B1 – สภาพดี", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่เหตุผล]"],
        ["B2 – ใช้งานปกติ", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่เหตุผล]"],
        ["B3 – เสื่อม", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่เหตุผล]"],
    ], [3.2, 3.5, 3.2, 3.8, 6.6])
    add_figure_placeholder(doc, "5.9", "ความสัมพันธ์ระหว่าง SoH จาก Quick Scan และ C10 reference", "Scatter plot: แกน x = C10 SoH (%), แกน y = Quick SoH (%), ใส่เส้น y=x และป้าย Battery ID. ถ้ามีข้อมูลไม่พอให้ใช้ grouped comparison chart แทน.", 4.2)

    add_heading(doc, "5.3.5 ความครบถ้วนของข้อมูลและ Cloud Traceability", level=2)
    add_table(doc, "ตารางที่ 5.11 การตรวจสอบ traceability ของข้อมูล", ["Run ID", "CSV metadata", "Mode", "ผลวิเคราะห์", "Cloud/dashboard", "รายงาน", "ผล"], [
        ["[ใส่ Run ID]", "profile/app version [ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
        ["[ใส่ Run ID]", "profile/app version [ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]", "[ใส่ผล]"],
    ], [2.2, 4.5, 2.0, 2.8, 3.2, 2.5, 2.0])
    add_figure_placeholder(doc, "5.10", "ตัวอย่างความต่อเนื่องของข้อมูลจาก CSV สู่ Cloud Dashboard", "ใส่ภาพ dashboard ของ Run ID เดียวกับ CSV ที่อ้างในตาราง 5.11 และใส่กรอบ/ลูกศรชี้ Battery ID, timestamp, test mode และผล grade. ปิดข้อมูลส่วนบุคคลหากมี.", 4.2)

    add_heading(doc, "5.4 ตารางตรวจสอบความสอดคล้องกับข้อกำหนด / Requirement Compliance Matrix")
    add_text(doc, "ตารางนี้ต้องสรุปผลทุก requirement รวมถึงข้อที่ไม่ผ่านและข้อที่ยังไม่มีหลักฐานเพียงพอ คำว่า PENDING ใช้เฉพาะกรณีที่มีแผนทดสอบชัดเจนแต่ยังทำไม่ครบ")
    add_table(doc, "ตารางที่ 5.12 ตารางตรวจสอบความสอดคล้องกับข้อกำหนด", ["Requirement", "Spec/Target", "Measured", "สถานะ", "หมายเหตุ"], [
        ["Transient sampling", "3–10 Hz", "Quick pulse/relax ≈3–6 Hz (เบื้องต้น)", "PRELIMINARY PASS", "ยืนยันใหม่ด้วย Mode"],
        ["Near-cutoff sampling", "3–10 Hz", "[ใส่ผลหลังแก้]", "PENDING", "ต้องเร่ง sampling ใกล้ 10.5 V"],
        ["DCIR repeatability", "R₀ CV ≤15%", "[ใส่ผล 3 run]", "PENDING", "ข้อมูล B007 เดิมมี 2 run"],
        ["SoC uncertainty", "±5%", "[ใส่ผลเทียบ reference]", "PENDING", "ต้องมี C10/controlled SoC"],
        ["SSR cutoff response", "≤1 s", "[ใส่ผล scope]", "PENDING", "ทดสอบ fault 3 รอบ"],
        ["OTP threshold", "45–60 °C", "[ใส่ผล]", "PENDING", "บันทึก sensor/reference"],
        ["Discharge cutoff", "10.4–10.6 V", "Quick 1: 10.18 V; Quick 2: 10.27 V", "FAIL", "ต้องแก้ near-cutoff logic และ retest"],
        ["Data traceability", "CSV → cloud → report", "[ใส่ผล]", "PENDING", "ใช้ Run ID เดียวกัน"],
    ], [3.7, 3.3, 4.5, 3.0, 4.3])

    add_heading(doc, "5.5 การทบทวนการออกแบบ / Design Review")
    add_text(doc, "ผลเบื้องต้นแสดงว่า Quick Scan สามารถเก็บข้อมูลความถี่สูงในช่วง mini-pulse และ relaxation ได้ ซึ่งเพียงพอสำหรับการวิเคราะห์ DCIR/ECM ในระดับเบื้องต้น อย่างไรก็ตาม CSV รุ่นเดิมยังไม่มี Mode จึงต้องอนุมาน phase จากเวลาและกระแส การเพิ่ม Mode, profile และ app version ใน CSV รุ่นใหม่จึงเป็นการปรับปรุงสำคัญต่อการตรวจสอบย้อนกลับของผล")
    add_text(doc, "ข้อจำกัดหลักที่พบจาก B007 คือ main discharge ใช้ low-rate logging ประมาณ 5 วินาทีต่อจุด ทำให้การยืนยัน cutoff ล่าช้าและแรงดันที่วัดได้ลดลงถึง 10.18–10.27 V ซึ่งต่ำกว่าเกณฑ์ 10.4–10.6 V ที่ระบุไว้ การปรับปรุงที่เสนอคือให้ระบบเพิ่ม sampling rate เป็น 3–10 Hz เมื่อแรงดันเข้าใกล้ cutoff และใช้ time-based debounce ที่สั้นพอ พร้อมบันทึกจุดเริ่มเตือน จุดสั่งตัด และแรงดันหลังปลดโหลด")
    add_text(doc, "ด้านการประเมินสุขภาพแบตเตอรี่ Quick Scan ควรรายงาน Quick Electrical Grade จาก DCIR/ECM แยกจาก Capacity Grade ที่ได้จาก C10 reference test แล้วกำหนด Verified Overall Grade จากผลที่อนุรักษ์นิยมกว่า วิธีนี้ป้องกันการตีความว่าแบตเตอรี่ได้เกรด A เพียงเพราะ parameter ของโมเดลหรือ Peukert factor เปลี่ยน โดยไม่มีการยืนยันความจุจริง")
    add_table(doc, "ตารางที่ 5.13 ประเด็นปรับปรุงจากผลทดสอบ", ["ประเด็น", "หลักฐาน", "แนวทางแก้", "การยืนยันหลังแก้"], [
        ["CSV เดิมไม่มี Mode", "แยก phase ได้เพียงการอนุมาน", "บันทึก Mode/profile/app version", "ตรวจ table 5.4 จาก CSV ใหม่"],
        ["Cutoff ต่ำกว่าเกณฑ์", "10.18 V และ 10.27 V", "near-cutoff high-rate + debounce", "ทดสอบ UVP 3 รอบ"],
        ["Quick grade ยังไม่ยืนยัน", "ไม่มี C10 reference ครบ", "แยก Electrical/Capacity/Overall Grade", "เปรียบเทียบ table 5.9–5.10"],
        ["Repeatability ยังไม่ครบ", "B007 มี 2 run", "รัน 3 รอบต่อแบตเตอรี่", "คำนวณ R₀ %CV"],
    ], [4.0, 4.0, 5.3, 5.8])

    add_text(doc, "รายการรูปที่ต้องจัดเตรียมก่อนส่ง: รูปที่ 5.1–5.10 ตาม placeholder ข้างต้น โดยใช้ข้อมูลจาก run ที่อ้างในตารางเดียวกันเท่านั้น รายละเอียด CSV เต็ม, raw waveform และผลรายรอบที่ยาวเกินหน้า ควรเก็บไว้ในภาคผนวกเพื่อให้บทนี้อ่านง่ายและตรวจสอบย้อนกลับได้", bold=True, before=8)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
