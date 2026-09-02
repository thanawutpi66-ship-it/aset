"""Replace the condensed Chapter 5 Table 5.2 with the detailed ABET draft table."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "output" / "docx" / "A19_capstone_design_ch5_revised.docx"
OUTPUT = ROOT / "output" / "docx" / "A19_capstone_design_ch5_tables_detailed.docx"
DRAFT = Path(r"C:\Users\boatl\Downloads\chapter_5_abet_detailed_draft.md")
TARGET_TABLE_INDEX = 10


def extract_markdown_table(markdown: str, caption: str) -> list[list[str]]:
    marker = markdown.index(caption)
    lines = markdown[marker:].splitlines()[1:]
    table_lines: list[str] = []
    for line in lines:
        if not line.strip():
            break
        if line.lstrip().startswith("|"):
            table_lines.append(line.strip())
    rows: list[list[str]] = []
    for line in table_lines:
        if re.match(r"^\|\s*-+", line):
            continue
        rows.append([cell.strip() for cell in line.strip("|").split("|")])
    if len(rows) != 16 or len(rows[0]) != 7:
        raise ValueError(f"Unexpected Table 5.2 shape: {len(rows)} x {len(rows[0]) if rows else 0}")
    return rows


def clear_and_add_text(cell, text: str, *, header: bool = False) -> None:
    bold = "**" in text
    clean = text.replace("**", "").replace("$", "").replace("<br>", "\n")
    cell.text = clean
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(7.5 if header else 8)
            run.bold = header or bold
            if header:
                run.font.color.rgb = RGBColor(255, 255, 255)


def set_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def prevent_row_split(row, repeat_header: bool = False) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    if repeat_header:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def make_table(doc: Document, rows: list[list[str]]):
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.12), Inches(0.98), Inches(0.78), Inches(0.78), Inches(0.83), Inches(0.72), Inches(0.82)]
    for col, width in zip(table.columns, widths):
        col.width = width
    for idx, value in enumerate(rows[0]):
        cell = table.rows[0].cells[idx]
        cell.width = widths[idx]
        set_shading(cell, "1F4E78")
        clear_and_add_text(cell, value, header=True)
    prevent_row_split(table.rows[0], repeat_header=True)
    for values in rows[1:]:
        cells = table.add_row().cells
        for idx, (cell, value) in enumerate(zip(cells, values)):
            cell.width = widths[idx]
            clear_and_add_text(cell, value)
        prevent_row_split(table.rows[-1])
    return table


def enable_field_updates(doc: Document) -> None:
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> None:
    if not SOURCE.exists() or not DRAFT.exists():
        raise FileNotFoundError("Missing source DOCX or Chapter 5 draft")
    rows = extract_markdown_table(DRAFT.read_text(encoding="utf-8"), "**ตารางที่ 5.2**")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    old_table = doc.tables[TARGET_TABLE_INDEX]
    new_table = make_table(doc, rows)
    old_table._tbl.addprevious(new_table._tbl)
    old_table._element.getparent().remove(old_table._element)
    enable_field_updates(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
