"""Summarize DOCX structure for a two-version content comparison."""

from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

from docx import Document


def clean(text: str) -> str:
    return " ".join(text.replace("\n", " ").split())


def summarize(path: Path) -> dict:
    doc = Document(path)
    paragraphs = [
        {
            "index": i,
            "style": para.style.name,
            "text": clean(para.text),
        }
        for i, para in enumerate(doc.paragraphs)
        if clean(para.text)
    ]
    tables = []
    for i, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([clean(cell.text)[:160] for cell in row.cells])
        tables.append(
            {
                "index": i,
                "rows": len(table.rows),
                "cols": len(table.columns),
                "preview": rows[:3],
            }
        )
    rels = Counter(rel.reltype.rsplit("/", 1)[-1] for rel in doc.part.rels.values())
    return {
        "name": path.name,
        "paragraph_count": len(doc.paragraphs),
        "nonempty_paragraph_count": len(paragraphs),
        "paragraph_style_counts": Counter(p["style"] for p in paragraphs),
        "headings": [p for p in paragraphs if p["style"].lower().startswith("heading")],
        "paragraphs": paragraphs,
        "tables": tables,
        "sections": len(doc.sections),
        "relationships": rels,
    }


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    payload = [summarize(Path(arg)) for arg in sys.argv[1:]]
    print(json.dumps(payload, ensure_ascii=False, indent=2, default=dict))


if __name__ == "__main__":
    main()
