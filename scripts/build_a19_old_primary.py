"""Build the consolidated A19 report with the 00_00 draft as its base."""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASE = Path(r"C:\Users\boatl\Downloads\A19_capstone_design_2026_00_00 (แก้แล้ว).docx")
UPDATE = Path(r"C:\Users\boatl\Downloads\A19_capstone_design_2026_07_24 (แก้แล้ว) (1).docx")
OUTPUT = Path(r"C:\Users\boatl\Projects\ASET_BATT\output\docx\A19_capstone_design_merged.docx")


def find_one(doc: Document, predicate, label: str):
    matches = [paragraph for paragraph in doc.paragraphs if predicate(paragraph.text.strip())]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one {label}, found {len(matches)}")
    return matches[0]


def body_elements_between(doc: Document, start, end) -> list:
    children = list(doc.element.body)
    begin = children.index(start._element)
    finish = children.index(end._element)
    elements = children[begin:finish]
    drawings = sum(len(element.findall('.//' + qn('a:blip'))) for element in elements)
    if drawings:
        raise RuntimeError("The requested merge block contains drawings that need relationship merging")
    return [deepcopy(element) for element in elements]


def insert_before(reference_element, elements: list) -> None:
    parent = reference_element.getparent()
    position = parent.index(reference_element)
    for offset, element in enumerate(elements):
        parent.insert(position + offset, element)


def remove_between(start, end) -> None:
    parent = start.getparent()
    children = list(parent)
    for child in children[children.index(start) : children.index(end)]:
        parent.remove(child)


def accept_tracked_changes(doc: Document) -> None:
    root = doc.element
    for deletion in list(root.iter(qn('w:del'))):
        deletion.getparent().remove(deletion)
    for insertion in list(root.iter(qn('w:ins'))):
        parent = insertion.getparent()
        position = parent.index(insertion)
        for child in list(insertion):
            insertion.remove(child)
            parent.insert(position, child)
            position += 1
        parent.remove(insertion)


def remove_red_review_notes(doc: Document) -> None:
    notes = []
    for color in doc.element.iter(qn('w:color')):
        if color.get(qn('w:val'), '').lower() != 'ff0000':
            continue
        paragraph = color
        while paragraph is not None and paragraph.tag != qn('w:p'):
            paragraph = paragraph.getparent()
        if paragraph is not None and '***' in ''.join(paragraph.itertext()):
            notes.append(paragraph)
    for paragraph in dict.fromkeys(notes):
        paragraph.getparent().remove(paragraph)


def request_field_update(doc: Document) -> None:
    settings = doc.settings.element
    update = settings.find(qn('w:updateFields'))
    if update is None:
        update = OxmlElement('w:updateFields')
        settings.append(update)
    update.set(qn('w:val'), 'true')


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(BASE, OUTPUT)
    base = Document(OUTPUT)
    update = Document(UPDATE)

    # Add the new verification chapter from the updated document before Ch. 6.
    update_ch5 = find_one(update, lambda text: text.startswith('บทที่ 5 '), 'updated Chapter 5')
    update_ch5_position = next(
        index
        for index, paragraph in enumerate(update.paragraphs)
        if paragraph.text.strip() == update_ch5.text.strip()
    )
    update_references_after_ch5 = next(
        paragraph
        for paragraph in update.paragraphs[update_ch5_position + 1 :]
        if paragraph.text.strip() == 'เอกสารอ้างอิง / References'
    )
    base_ch6 = find_one(base, lambda text: text.startswith('บทที่ 6 '), 'base Chapter 6')
    insert_before(
        base_ch6._element,
        body_elements_between(update, update_ch5, update_references_after_ch5),
    )

    # Replace the older short Ch. 8 with the detailed, non-duplicated Ch. 8.
    update_ch8 = find_one(
        update,
        lambda text: text.startswith('บทที่ 8 ') and 'Chapter 8:' in text,
        'detailed updated Chapter 8',
    )
    update_duplicate_ch8 = find_one(update, lambda text: text == 'บทที่ 8 บทสรุปและข้อเสนอแนะ', 'duplicate Chapter 8')
    detailed_ch8 = body_elements_between(update, update_ch8, update_duplicate_ch8)

    base_ch8 = find_one(base, lambda text: text.startswith('บทที่ 8 '), 'base Chapter 8')
    base_references = next(
        paragraph
        for paragraph in reversed(base.paragraphs)
        if paragraph.text.strip() == 'เอกสารอ้างอิง / References'
    )
    remove_between(base_ch8._element, base_references._element)
    insert_before(base_references._element, detailed_ch8)

    accept_tracked_changes(base)
    remove_red_review_notes(base)
    request_field_update(base)
    base.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
