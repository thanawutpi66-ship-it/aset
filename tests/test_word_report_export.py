"""Smoke test for the editable Chapter-5 experiment evidence report."""
import csv
import os
import tempfile
import unittest
from types import SimpleNamespace

from docx import Document

from aset_batt.storage.word_report import generate_word_report


def _config():
    return SimpleNamespace(battery=SimpleNamespace(
        battery_type="LeadAcid", product_name="B007", cells_series=6,
        cells_parallel=1, pack_nominal_voltage=12.0, rated_capacity=5.3,
    ))


class TestWordExperimentReport(unittest.TestCase):
    def test_export_contains_report_ready_evidence_sections(self):
        with tempfile.TemporaryDirectory() as tmp:
            csv_path = os.path.join(tmp, "session.csv")
            with open(csv_path, "w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=[
                    "Elapsed_s", "Voltage_V", "Current_A", "SoC_pct", "Temperature_C", "Mode"])
                writer.writeheader()
                writer.writerows([
                    {"Elapsed_s": 0, "Voltage_V": 12.7, "Current_A": 0, "SoC_pct": 100, "Temperature_C": 25, "Mode": "REST"},
                    {"Elapsed_s": 1, "Voltage_V": 12.4, "Current_A": 5.3, "SoC_pct": 99, "Temperature_C": 25, "Mode": "MINI_PULSE"},
                    {"Elapsed_s": 2, "Voltage_V": 12.6, "Current_A": 0, "SoC_pct": 99, "Temperature_C": 25, "Mode": "RELAX"},
                    {"Elapsed_s": 3, "Voltage_V": 12.3, "Current_A": 0.53, "SoC_pct": 99, "Temperature_C": 25, "Mode": "MAIN_DISCHARGE"},
                ])
            output = os.path.join(tmp, "evidence.docx")
            generate_word_report(output, _config(), analysis={
                "overall_grade": "REVIEW", "capacity_grade": "REVIEW", "electrical_grade": "A",
                "capacity_ah": 3.2, "capacity_rate_normalized_ah": 4.0, "capacity_basis": "MAIN_DISCHARGE",
                "soh": 60.0, "soh_basis": "rate/SoC-normalised estimate; not a full-charge measurement",
                "dcir_mohm": 40.0, "dcir_n_steps": 1, "r0_mohm": 30.0, "r1_mohm": 10.0,
                "ecm_r2": 0.95, "quality_warnings": ["capacity grade withheld — not a C10 run"],
            }, csv_path=csv_path)
            self.assertTrue(os.path.exists(output))
            document = Document(output)
            text = "\n".join(p.text for p in document.paragraphs)
            table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            self.assertIn("ASET Battery Experiment Evidence Report", text)
            self.assertIn("Chapter 5 Evidence Checklist", text)
            self.assertIn("Report-ready interpretation", text)
            self.assertIn("Verified Overall Grade", table_text)
            self.assertIn("capacity grade withheld", table_text)


if __name__ == "__main__":
    unittest.main()
