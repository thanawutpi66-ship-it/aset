"""Export one traceable battery-test session as a Word evidence report.

The document is deliberately evidence-gated: it reports a capacity estimate,
electrical result, and verified overall grade separately.  It never turns a
Quick Scan or an incomplete CSV into a claimed C10 capacity acceptance.
"""
from __future__ import annotations

import csv
import json
import logging
import math
import os
import tempfile
from datetime import datetime
from pathlib import Path

import numpy as np

logger = logging.getLogger(__name__)

_BLUE = "2E74B5"
_DARK_BLUE = "1F4D78"
_HEADER_FILL = "E8EEF5"
_CALLOUT_FILL = "F4F6F9"
_TEXT = "0B2545"
_PAGE_WIDTH_DXA = 9360


def _safe_float(value, default=float("nan")):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _fmt(value, digits=2, unit="", missing="N/A"):
    value = _safe_float(value)
    if not math.isfinite(value):
        return missing
    return f"{value:.{digits}f}{unit}"


def _read_session(csv_path: str) -> dict:
    data = {"t": [], "v": [], "i": [], "soc": [], "temp": [], "mode": []}
    if not csv_path or not os.path.exists(csv_path):
        return {key: np.asarray([], float) if key != "mode" else [] for key in data}
    with open(csv_path, "r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(line for line in handle if not line.lstrip().startswith("#")):
            data["t"].append(_safe_float(row.get("Elapsed_s")))
            data["v"].append(_safe_float(row.get("Voltage_V")))
            data["i"].append(_safe_float(row.get("Current_A")))
            data["soc"].append(_safe_float(row.get("SoC_pct")))
            data["temp"].append(_safe_float(row.get("Temperature_C")))
            data["mode"].append(str(row.get("Mode") or "").strip())
    for key in ("t", "v", "i", "soc", "temp"):
        data[key] = np.asarray(data[key], float)
    return data


def _session_stats(data: dict) -> dict:
    t = data["t"]
    finite_t = t[np.isfinite(t)]
    dt = np.diff(finite_t)
    dt = dt[np.isfinite(dt) & (dt > 0)]
    modes = sorted({m for m in data["mode"] if m})
    return {
        "samples": int(len(t)),
        "duration_s": float(np.nanmax(finite_t)) if finite_t.size else float("nan"),
        "median_dt_s": float(np.median(dt)) if dt.size else float("nan"),
        "median_hz": float(1.0 / np.median(dt)) if dt.size and np.median(dt) > 0 else float("nan"),
        "modes": modes,
    }


def _render_figures(data: dict, directory: str) -> list[tuple[str, str]]:
    """Create report figures; a missing matplotlib is non-fatal to export."""
    t = data["t"]
    if t.size < 2:
        return []
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception as exc:
        logger.warning("Word report graphs unavailable: %s", exc)
        return []

    out = []
    finite = np.isfinite(t)
    if not finite.any():
        return out
    x = t / 60.0
    fig, (ax_v, ax_i) = plt.subplots(2, 1, figsize=(7.0, 4.4), dpi=160, sharex=True)
    ax_v.plot(x, data["v"], color="#005A9E", linewidth=1.2)
    ax_v.set_ylabel("Voltage (V)")
    ax_v.grid(alpha=0.25)
    ax_i.plot(x, data["i"], color="#D83B01", linewidth=1.0)
    ax_i.axhline(0.0, color="#666666", linewidth=0.6)
    ax_i.set_xlabel("Elapsed time (min)")
    ax_i.set_ylabel("Current (A)\n(discharge +)")
    ax_i.grid(alpha=0.25)
    fig.tight_layout()
    vi_path = os.path.join(directory, "voltage_current.png")
    fig.savefig(vi_path, bbox_inches="tight")
    plt.close(fig)
    out.append(("Figure 1. Voltage and current profile from the raw CSV.", vi_path))

    if np.isfinite(data["soc"]).any() or np.isfinite(data["temp"]).any():
        fig, ax = plt.subplots(figsize=(7.0, 3.4), dpi=160)
        if np.isfinite(data["soc"]).any():
            ax.plot(x, data["soc"], color="#2E8B57", linewidth=1.2, label="SoC")
            ax.set_ylabel("SoC (%)", color="#2E8B57")
            ax.set_ylim(-2, 102)
        ax2 = ax.twinx()
        if np.isfinite(data["temp"]).any():
            ax2.plot(x, data["temp"], color="#B22222", linewidth=1.0, label="Temperature")
            ax2.set_ylabel("Temperature (°C)", color="#B22222")
        ax.set_xlabel("Elapsed time (min)")
        ax.grid(alpha=0.25)
        fig.tight_layout()
        state_path = os.path.join(directory, "soc_temperature.png")
        fig.savefig(state_path, bbox_inches="tight")
        plt.close(fig)
        out.append(("Figure 2. Logged SoC and temperature trend.", state_path))
    return out


def _set_cell_fill(cell, fill: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def _set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths_dxa: list[int]):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Twips
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    total = sum(widths_dxa)
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    indent = tbl_pr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths_dxa):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            cell.width = Twips(width)
            _set_cell_margins(cell)


def _write_cell(cell, text, *, bold=False, color=_TEXT, center=False):
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = cell.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)


def _add_table(doc, headers: list[str], rows: list[list[str]], widths_dxa: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    _set_table_geometry(table, widths_dxa)
    for cell, text in zip(table.rows[0].cells, headers):
        _set_cell_fill(cell, _HEADER_FILL)
        _write_cell(cell, text, bold=True, color=_DARK_BLUE, center=True)
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            _write_cell(cell, text)
    return table


def _add_heading(doc, text: str, level: int):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def _analysis_rows(analysis: dict) -> list[list[str]]:
    return [
        ["Quick Scan Grade", str(analysis.get("quick_grade", "REVIEW")),
         str(analysis.get("quick_grade_basis", "Not a Quick Scan record."))],
        ["Verified Overall Grade", str(analysis.get("overall_grade", analysis.get("grade", "REVIEW"))),
         "Valid only when both capacity and electrical evidence are available."],
        ["Capacity Grade", str(analysis.get("capacity_grade", "REVIEW")),
         "Requires phase-labelled, full C-rate reference discharge to cut-off."],
        ["Electrical Grade", str(analysis.get("electrical_grade", "REVIEW")),
         "Requires valid DCIR/ECM pulse data."],
        ["Observed capacity", _fmt(analysis.get("capacity_ah"), 3, " Ah"),
         str(analysis.get("capacity_basis", "not available"))],
        ["Rate-normalised capacity", _fmt(analysis.get("capacity_rate_normalized_ah", analysis.get("capacity_norm_ah")), 3, " Ah"),
         "Diagnostic estimate; not capacity acceptance by itself."],
        ["Observed capacity fraction", _fmt(analysis.get("soh"), 1, " %"),
         str(analysis.get("soh_basis", ""))],
        ["Peukert-corrected SoH", _fmt(analysis.get("soh_est"), 1, " %"),
         "Used for Quick Scan Grade; not a measured C10 capacity result."],
        ["DCIR @ edge", _fmt(analysis.get("dcir_mohm"), 2, " mΩ"),
         f"valid steps: {analysis.get('dcir_n_steps', 0)}"],
        ["ECM R0 / R1", f"{_fmt(analysis.get('r0_mohm'), 2, ' mΩ')} / {_fmt(analysis.get('r1_mohm'), 2, ' mΩ')}",
         f"R²: {_fmt(analysis.get('ecm_r2'), 3)}"],
    ]


def generate_word_report(path, config, estimator=None, analysis=None, csv_path=None):
    """Create an editable Word report from one test session.

    It is suitable as a Chapter 5 evidence appendix: raw-data provenance,
    calculated metrics, quality limits, result tables, and figures are embedded
    in one `.docx`.  Missing evidence remains explicitly ``REVIEW``/``N/A``.
    """
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml.ns import qn

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in (("Heading 1", 16, _BLUE, 16, 8),
                                              ("Heading 2", 13, _BLUE, 12, 6),
                                              ("Heading 3", 12, _DARK_BLUE, 8, 4)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = "ASET BATTERY TESTER | EXPERIMENT EVIDENCE REPORT"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("6B7280")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Generated from the recorded CSV session | ")
    footer.add_run(datetime.now().strftime("%Y-%m-%d %H:%M"))
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("6B7280")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("ASET Battery Experiment Evidence Report")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor.from_string(_DARK_BLUE)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Editable Word appendix for Chapter 5: Verification and Validation").italic = True
    subtitle.runs[0].font.color.rgb = RGBColor.from_string("4B5563")

    b = config.battery
    data = _read_session(csv_path)
    stats = _session_stats(data)
    analysis_dict = analysis if isinstance(analysis, dict) else {}
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    _add_heading(doc, "1. Test and Data Provenance", 1)
    meta = {}
    if csv_path and os.path.exists(csv_path + ".meta.json"):
        try:
            meta = json.loads(Path(csv_path + ".meta.json").read_text(encoding="utf-8"))
        except (OSError, ValueError):
            meta = {}
    _add_table(doc, ["Field", "Recorded value"], [
        ["Battery under test", str(getattr(b, "product_name", "") or b.battery_type)],
        ["Chemistry / configuration", f"{b.battery_type}; {b.cells_series}S{b.cells_parallel}P"],
        ["Nominal voltage / rated capacity", f"{b.pack_nominal_voltage:.2f} V / {b.rated_capacity:.2f} Ah"],
        ["Raw CSV", os.path.basename(csv_path) if csv_path else "Not available"],
        ["Session ID / CSV SHA-256", f"{meta.get('session_id', 'Not recorded')} / {meta.get('sha256', 'Not finalized')}"],
        ["Protocol / analysis version", f"{(meta.get('protocol') or {}).get('id', 'Not recorded')} / {(meta.get('protocol') or {}).get('analysis_version', 'Not recorded')}"],
        ["Operator / software version", f"{meta.get('operator', 'Not recorded')} / {meta.get('app_version', 'Not recorded')}"],
        ["Samples / duration", f"{stats['samples']} / {_fmt(stats['duration_s'] / 60.0, 2, ' min')}"],
        ["Sampling interval / rate", f"{_fmt(stats['median_dt_s'], 3, ' s')} / {_fmt(stats['median_hz'], 2, ' Hz')}"],
        ["Recorded phases", ", ".join(stats["modes"]) if stats["modes"] else "Not recorded (legacy CSV)"],
    ], [2700, 6660])

    _add_heading(doc, "2. Analysis and Verification Summary", 1)
    if analysis_dict:
        _add_table(doc, ["Metric", "Measured result", "Interpretation"],
                   _analysis_rows(analysis_dict), [2550, 2200, 4610])
        narrative = doc.add_paragraph()
        narrative.add_run("Report-ready interpretation: ").bold = True
        overall = analysis_dict.get("overall_grade", analysis_dict.get("grade", "REVIEW"))
        narrative.add_run(
            f"This session produced Verified Overall Grade {overall}. "
            f"Capacity Grade is {analysis_dict.get('capacity_grade', 'REVIEW')} and Electrical Grade is "
            f"{analysis_dict.get('electrical_grade', 'REVIEW')}. "
            "A REVIEW result means the missing or invalid evidence must be completed before claiming compliance.")
    else:
        doc.add_paragraph("No unified analysis result was supplied; this report preserves raw-data evidence only.")

    _add_heading(doc, "3. Data-quality Evidence and Limitations", 1)
    warnings = analysis_dict.get("quality_warnings") or []
    if warnings:
        _add_table(doc, ["No.", "Recorded quality finding"],
                   [[str(idx), str(w)] for idx, w in enumerate(warnings, 1)], [700, 8660])
    else:
        note = doc.add_table(rows=1, cols=1)
        _set_table_geometry(note, [_PAGE_WIDTH_DXA])
        _set_cell_fill(note.cell(0, 0), _CALLOUT_FILL)
        _write_cell(note.cell(0, 0), "No analysis warnings were supplied. This does not prove calibration, safety-interlock, or cloud verification; those require their own test evidence.")

    _add_heading(doc, "4. Figures from the Recorded Session", 1)
    with tempfile.TemporaryDirectory(prefix="aset_word_report_") as temp_dir:
        figures = _render_figures(data, temp_dir)
        if figures:
            for caption, figure_path in figures:
                doc.add_picture(figure_path, width=Inches(6.35))
                p = doc.add_paragraph(caption)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.runs[0].italic = True
                p.runs[0].font.size = Pt(9)
        else:
            doc.add_paragraph("Graphs could not be generated because the CSV has insufficient rows or the chart dependency is unavailable.")

        _add_heading(doc, "5. Chapter 5 Evidence Checklist", 1)
        _add_table(doc, ["Chapter 5 evidence", "This export provides", "Still requires separate test"], [
            ["Measurement and logging", "CSV provenance, sample count/rate, V/I/SoC/temperature figures", "Instrument calibration / 4-wire accuracy"],
            ["Sequence automation", "Phase labels and end-to-end session trace when logged", "Repeat-run acceptance statistics"],
            ["Battery algorithm", "DCIR, ECM, capacity, grade split, quality warnings", "HPPC/C10 repeatability comparison"],
            ["Safety and interlock", "Session warning record only", "UVP/OVP/OTP/E-stop trip-time evidence"],
            ["Data and cloud", "CSV filename and metadata sidecar", "Dashboard capture / cloud-record audit"],
        ], [2400, 3600, 3360])

        doc.add_paragraph(
            "Use this file as a test-evidence appendix or copy its tables/figures into Chapter 5. "
            "Do not replace the report's requirement-compliance matrix with a single session result.")
        os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
        doc.save(path)
    logger.info("Word experiment report written: %s", path)
    return path
